import json

import openpyxl
import pytest

from src.domain.file_rows import row_key
from src.infrastructure.files.table_file import TableFileError
from src.infrastructure.recipients.file_recipient_source import FileRecipientSource


def write_csv(path, text: str) -> str:
    """Write a UTF-8 .csv file and return its path as text."""
    path.write_bytes(text.encode("utf-8"))
    return str(path)


def source_for(path: str, sheet: str | None = None) -> FileRecipientSource:
    """Build a source reading the 'Correo' and 'Adjunto' columns of *path*."""
    return FileRecipientSource(path, sheet, "Correo", "Adjunto")


class TestFileRecipientSourceRows:
    """Covers turning the rows of a file into recipients."""

    def test_maps_a_row_to_a_recipient(self, tmp_path):
        """A row becomes a recipient with its stripped email, attachment, row number and key."""
        path = write_csv(tmp_path / "f.csv", "Correo;Adjunto\n a@x.com ;factura.pdf\n")
        [recipient] = source_for(path).get_recipients()
        assert (recipient.email, recipient.attachment, recipient.row) == ("a@x.com", "factura.pdf", 2)
        assert recipient.key == row_key("a@x.com", "factura.pdf")
        assert recipient.skip_reason is None

    def test_empty_and_invalid_emails_are_skipped(self, tmp_path):
        """Rows with an empty or invalid address carry the no_email and invalid_email reasons."""
        path = write_csv(tmp_path / "f.csv", "Correo;Adjunto;Otra\n;a.pdf;x\na@x.com;b@x.com;b.pdf\nc@x.com, d@x.com;c.pdf\n")
        reasons = [r.skip_reason for r in source_for(path).get_recipients()]
        assert reasons == ["no_email", None, "invalid_email"]

    def test_rows_without_attachment_are_skipped(self, tmp_path):
        """A row with an address but no attachment carries the no_attachment reason."""
        path = write_csv(tmp_path / "f.csv", "Correo;Adjunto\na@x.com;\n")
        assert source_for(path).get_recipients()[0].skip_reason == "no_attachment"

    def test_repeated_rows_are_skipped_as_duplicates(self, tmp_path):
        """A second row with the same address (any case) and attachment is a duplicate."""
        path = write_csv(tmp_path / "f.csv", "Correo;Adjunto\na@x.com;a.pdf\nA@X.com;a.pdf\na@x.com;b.pdf\n")
        reasons = [r.skip_reason for r in source_for(path).get_recipients()]
        assert reasons == [None, "duplicate_row", None]

    def test_numeric_attachment_cells_lose_their_decimals(self, tmp_path):
        """An attachment written as a whole number in Excel is read without decimals."""
        workbook = openpyxl.Workbook()
        workbook.active.append(["Correo", "Adjunto"])
        workbook.active.append(["a@x.com", 1234.0])
        workbook.save(tmp_path / "f.xlsx")
        [recipient] = source_for(str(tmp_path / "f.xlsx")).get_recipients()
        assert recipient.attachment == "1234"

    def test_reads_the_chosen_sheet(self, tmp_path):
        """With a sheet given, the rows come from that sheet."""
        workbook = openpyxl.Workbook()
        workbook.active.title = "Octubre"
        workbook.active.append(["Correo", "Adjunto"])
        workbook.active.append(["a@x.com", "a.pdf"])
        november = workbook.create_sheet("Noviembre")
        november.append(["Correo", "Adjunto"])
        november.append(["b@x.com", "b.pdf"])
        workbook.save(tmp_path / "f.xlsx")
        [recipient] = source_for(str(tmp_path / "f.xlsx"), "Noviembre").get_recipients()
        assert recipient.email == "b@x.com"

    def test_file_is_read_again_on_every_call(self, tmp_path):
        """Changes saved to the file between the summary and the send are picked up."""
        path = write_csv(tmp_path / "f.csv", "Correo;Adjunto\na@x.com;a.pdf\n")
        source = source_for(path)
        source.get_recipients()
        write_csv(tmp_path / "f.csv", "Correo;Adjunto\na@x.com;a.pdf\nb@x.com;b.pdf\n")
        assert len(source.get_recipients()) == 2

    def test_missing_column_is_reported(self, tmp_path):
        """A chosen column that is no longer in the file is reported with its name."""
        path = write_csv(tmp_path / "f.csv", "Email;Adjunto\na@x.com;a.pdf\n")
        with pytest.raises(TableFileError) as info:
            source_for(path).get_recipients()
        assert (info.value.code, info.value.details) == ("missing_column", {"name": "Correo"})


class TestFileRecipientSourceIdentity:
    """Covers how a file source is identified in a campaign and in the log."""

    def test_identity_uses_the_file_name_without_its_folder(self, tmp_path):
        """Moving the file to another folder keeps the same identity, so a send can be resumed."""
        (tmp_path / "a").mkdir()
        (tmp_path / "b").mkdir()
        first = FileRecipientSource(str(tmp_path / "a" / "clientes.csv"), None, "Correo", "Adjunto")
        second = FileRecipientSource(str(tmp_path / "b" / "clientes.csv"), None, "Correo", "Adjunto")
        assert first.identity == second.identity

    def test_identity_includes_sheet_and_email_column(self, tmp_path):
        """Another sheet or another email column makes a different campaign."""
        path = str(tmp_path / "clientes.xlsx")
        base = FileRecipientSource(path, "Octubre", "Correo", "Adjunto").identity
        assert FileRecipientSource(path, "Noviembre", "Correo", "Adjunto").identity != base
        assert FileRecipientSource(path, "Octubre", "Email", "Adjunto").identity != base

    def test_identity_can_never_match_an_agenda_group(self, tmp_path):
        """The identity names the kind of source, so it never equals a group ID."""
        identity = FileRecipientSource(str(tmp_path / "123"), None, "Correo", "Adjunto").identity
        assert json.loads(identity)[0] == "file"

    def test_attachment_field_is_the_attachment_column(self, tmp_path):
        """The campaign's attachment field is the chosen attachment column."""
        assert source_for(str(tmp_path / "f.csv")).attachment_field == "Adjunto"

    def test_log_fields_name_the_file_sheet_and_email_column(self, tmp_path):
        """The log names the file (without its folder), the sheet and the email column."""
        source = FileRecipientSource(str(tmp_path / "clientes.xlsx"), "Hoja1", "Correo", "Adjunto")
        assert source.log_fields == {"file": "clientes.xlsx", "sheet": "Hoja1", "email_column": "Correo"}

    def test_log_fields_leave_out_the_sheet_of_a_csv(self, tmp_path):
        """A CSV file has no sheet to name."""
        assert "sheet" not in source_for(str(tmp_path / "f.csv")).log_fields
